from pptx import Presentation

prs = Presentation("/Users/vcv/Downloads/W4_KNN_Bayes_2026 (1).pptx")

for i, slide in enumerate(prs.slides, start=1):
    print(f"\n--- Slide {i} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            print(shape.text)